"""按用户给定个人模板导出岗位xlsx，根目录汇总，独立审核录入及材料索引。"""
import re
from collections import defaultdict
from pathlib import Path
import xlrd
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from excel_links import link_person,save_cached

ROOT=Path(__file__).resolve().parent


def suggested(score):
    """人工复核分优先；未审核时展示填报试算，真实0分不回退。"""
    return score['approved'] if score.get('approved') is not None else score.get('tentative')


def display_status(status):
    """核对一致已写在赋分依据中，界面和导出省去重复状态。"""
    return '' if status=='材料信息一致（建议）' else status


def clean_reason(value):
    """用户表只保留业务说明，去掉内部标准位置及网址。"""
    lines=[line for line in str(value or '').splitlines() if not line.strip().startswith('标准：')]
    cleaned=re.sub(r'https?://\S+','', '\n'.join(lines))
    cleaned=re.sub(r'[；;]?学校标识码\d+[；;]?','',cleaned)
    return cleaned.strip('；; \n')


def brief_info(person, codes):
    """按项目展示必要信息，技能项不重复拼接整段实习等共享自由文本。"""
    p=person['fields']
    mapping={'S01':['最高学历毕业院校'],'S02':['最高学历'],'S03':['最高学历学习形式'],
             'S04':['政治面貌'],'S05':['最高学历毕业时间'],'S06':[k for k in p if k.startswith('在读期间有无')],
             'S07':['奖学金'],'S08':['表彰奖励'],'S09':['计算机等级'],'S10':['英语等级'],
             'S13':['性别']}
    pieces=[]
    for code in codes:
        if code in mapping:
            for field in mapping[code]:
                if p.get(field) not in (None,''):pieces.append(str(p[field]))
        elif code=='S11':
            pieces.append('其他证书：见凭证核验状态')
        elif code in {'S14','S15'}:
            score=next(s for s in person['scores'] if s['code']==code)
            pieces.append(score.get('key_info','见赋分原因'))
        else:
            score=next(s for s in person['scores'] if s['code']==code)
            pieces.append(score.get('raw',''))
    return '\n'.join(dict.fromkeys(x for x in pieces if x))
HEADERS=['人员编号','岗位编码','姓名','项目编号','项目','输入指纹','当前结论','填报试算分','材料复核分',
         '复核结论','复核得分','复核依据','审核人','证明材料路径（每行一个）','已核对原图','审核阶段',
         '境外学位取得日期','境外认证完成日期','审核时间']


def safe(value):
    if isinstance(value,str) and value.lstrip().startswith(('=','+','-','@')):return "'"+value
    return value


def table(book,name,headers,rows):
    sheet=book.create_sheet(name)
    for row in [headers]+rows:sheet.append([safe(x) for x in row])
    sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
    for cell in sheet[1]:cell.font=Font(name='微软雅黑',bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='24476A')
    for row in sheet:
        for c in row:c.alignment=Alignment(vertical='top',wrap_text=True)
    for col in range(1,len(headers)+1):sheet.column_dimensions[get_column_letter(col)].width=22 if col<5 else 35
    sheet.row_dimensions[1].height=36
    return sheet


def template(book,name,master):
    original=xlrd.open_workbook(ROOT/'data/templates'/('硕士赋分模板.xls' if master else '本科赋分模板.xls'),formatting_info=True)
    src=original.sheet_by_index(0);dest=book.create_sheet(name)
    edge=Side(style='thin',color='444444')
    for row in range(src.nrows):
        dest.row_dimensions[row+1].height=src.rowinfo_map[row].height/20 if row in src.rowinfo_map else 30
        for col in range(7):
            value=src.cell_value(row,col)
            if row>=2 and col>=3:value=None  # 模板示例个人信息和示例分数全部清空。
            cell=dest.cell(row+1,col+1,value if value!='' else None)
            xf=original.xf_list[src.cell_xf_index(row,col)];font=original.font_list[xf.font_index]
            cell.font=Font(name=font.name,size=font.height/20,bold=bool(font.bold))
            cell.alignment=Alignment(horizontal='left' if col in (3,5,6) else 'center',vertical='center',wrap_text=True)
            cell.border=Border(left=edge,right=edge,top=edge,bottom=edge)
    for c in range(7):dest.column_dimensions[get_column_letter(c+1)].width=src.colinfo_map[c].width/256 if c in src.colinfo_map else 20
    for r1,r2,c1,c2 in src.merged_cells:dest.merge_cells(start_row=r1+1,end_row=r2,start_column=c1+1,end_column=c2)
    dest.sheet_properties.pageSetUpPr.fitToPage=True
    dest.page_setup.orientation='landscape';dest.page_setup.paperSize=dest.PAPERSIZE_A4
    dest.page_setup.fitToWidth=1;dest.page_setup.fitToHeight=0
    dest.print_options.horizontalCentered=True;dest.print_area=f'A1:G{src.nrows}'
    return dest


def export_batch(report,out):
    """分别生成岗位汇总与个人工作簿；按入口配置选择是否导出复核录入表。"""
    out=Path(out);groups=defaultdict(list);reviews=[]
    for person in report['persons']:
        groups[(person['job'],person['unit'],person['job_name'],person['master'])].append(person)
        for item in person['qualification']+person['scores']:
            d=report['decisions'].get('|'.join((person['pid'],person['job'],item['code'])),{})
            reviews.append([person['pid'],person['job'],person['name'],item['code'],item['name'],person['fingerprint'],
                item['status'],item.get('tentative'),item.get('approved'),d.get('conclusion',''),d.get('score'),
                d.get('reason',''),d.get('reviewer',''),'\n'.join(d.get('proofs',[])),'是' if d.get('proof_checked') else '',
                d.get('stage','初审'),d.get('degree_date',''),d.get('certification_date',''),d.get('reviewed_at','')])
    internal=out/'内部记录';internal.mkdir(parents=True,exist_ok=True)
    if report.get('options',{}).get('export_review',True):
        forms=openpyxl.Workbook();forms.remove(forms.active);f=table(forms,'复核录入',HEADERS,reviews)
        for col,choices in [('J','符合,不符合,待复核,不适用'),('O','是,否'),('P','初审,复审')]:
            dv=DataValidation(type='list',formula1=f'"{choices}"');f.add_data_validation(dv)
            if f.max_row>1:dv.add(f'{col}2:{col}{f.max_row}')
        forms.save(internal/'人工复核录入.xlsx')
    for (job,unit,job_name,master),persons in groups.items():
        filename=re.sub(r'[<>:"/\\|?*]','_',f'{job}_{unit}_{job_name}_{"硕士" if master else "本科"}')[:130]+'.xlsx'
        book=openpyxl.Workbook();book.remove(book.active);summary=[];summary_cache={}
        for n,person in enumerate(persons,1):
            eligible=person['conclusion'] in {'符合','待复核'}
            scores=person['scores']
            unknown=sum(suggested(s) is None for s in scores)
            pending=sum(s['approved'] is None and s.get('status')!='按本人填报' for s in scores)
            issues=[q['name']+'：'+clean_reason(q['reason']) for q in person['qualification'] if q['status'] in {'不符合','待复核'}]
            if report.get('options',{}).get('ocr_scope')=='education_party':issues.insert(0,'本次仅试跑前四项材料核验，其他项目尚未进行完整材料核验。')
            if report.get('options',{}).get('ocr_scope')=='education_party_scholarship':issues.insert(0,'本次试跑学历、政治面貌、毕业年限及奖学金材料，其他项目尚未进行完整材料核验。')
            if report.get('options',{}).get('ocr_scope')=='skills':issues.insert(0,'本次仅试跑技能等级三项（计算机、英语、其他证书），其他项目尚未进行完整材料核验。')
            if report.get('options',{}).get('ocr_scope')=='management':issues.insert(0,'本次仅试跑学生工作及科研/实习实践材料，其他项目尚未完整核验。')
            stale=[x for x in report.get('stale_reviews',[]) if x['key'].startswith(person['pid']+'|'+person['job']+'|')]
            if stale:issues.append(f'{len(stale)}项旧审核因输入变化需要重新确认')
            if eligible:
                total=sum(suggested(s) or 0 for s in scores)
                state=f'建议分小计；{unknown}项分数待定，{pending}项材料待核验'
            else:total=None;state='不赋分' if person['conclusion']=='不符合' else '未处理'
            summary.append([str(person['ident']),person['name'],person['category']['name'],person['conclusion'],total,state,'\n'.join(issues)])
            person['output_file']=str((out/filename).resolve())
        summary_sheet=table(book,'岗位汇总',['身份证号','姓名','考生类别','资格建议结论','建议分小计','核验状态','说明'],summary)
        summary_sheet.column_dimensions['A'].width=24
        for row in range(2,summary_sheet.max_row+1):summary_sheet.cell(row,1).number_format='@'
        summary_sheet.column_dimensions['G'].width=70
        for n,person in enumerate(persons,1):
            if person['conclusion'] not in {'符合','待复核'}:continue
            base_name=re.sub(r'[\\/*?:\[\]]','_',person['name']).strip("'") or '考生'
            name=base_name[:26];suffix=2
            while name in book.sheetnames:name=base_name[:24]+'_'+str(suffix);suffix+=1
            personal=openpyxl.Workbook();personal.remove(personal.active)
            sheet=template(personal,'个人赋分表',person['master'])
            personal_dir=out/'个人赋分表'/Path(filename).stem
            personal_dir.mkdir(parents=True,exist_ok=True)
            personal_path=personal_dir/(f'{person["ident"]}_{name}.xlsx')
            summary_sheet.cell(n+1,2).hyperlink=personal_path.resolve().as_uri()
            sheet['A1']=f'通过资格审查人员赋分表（建议稿）\n姓名：{person["name"]}  身份证号：{person["ident"]}\n资格状态：{person["conclusion"]}'
            sheet.row_dimensions[1].height=65;sheet['E2']='建议得分'
            rows=[['S01'],['S02'],['S03'],['S04'],['S05'],['S06'],['S07'],['S08'],['S09','S10','S11']]
            if person['master']:rows.append(['S12'])
            rows += [['S13'],['S14'],['S15']]
            scores={s['code']:s for s in person['scores']}
            for row,codes in enumerate(rows,3):
                sheet.cell(row,1,row-2)
                values=[scores[c] for c in codes]
                sheet.cell(row,4,safe(brief_info(person,codes)))
                known=[suggested(s) for s in values if suggested(s) is not None]
                if known:sheet.cell(row,5,sum(known))
                reasons=[];evidence=[]
                for score in values:
                    value=suggested(score)
                    label=score['name']+'：' if len(values)>1 else ''
                    status=display_status(score['status'])
                    reasons.append(label+('待定' if value is None else str(value)+'分')+('；'+status if status else '')+'\n'+clean_reason(score['reason']))
                    evidence.append('本人填报' if score['code'] in {'S12','S13'} else label+('、'.join(Path(p).name for p in score['proofs']) or ('无最高学历成绩单' if score['code']=='S06' and score['approved'] is None else '按明确填报不加分' if score['approved']==0 else '凭证待核对')))
                if len(known)<len(values):reasons.insert(0,'本行仅计已知建议分，其他子项待定。')
                sheet.cell(row,6,safe('\n'.join(reasons)));sheet.cell(row,7,safe('\n'.join(evidence)))
                links=[p for v in values for p in v['proofs']]
                if links:sheet.cell(row,7).hyperlink=Path(links[0]).as_uri()
                sheet.row_dimensions[row].height=max(sheet.row_dimensions[row].height or 30,75)
            total_row=3+len(rows)
            sheet.cell(total_row,1,'建议分小计')
            total=sum(suggested(s) or 0 for s in person['scores'])
            sheet.cell(total_row,5,f'=SUM(E3:E{total_row-1})')
            unknown=sum(suggested(s) is None for s in person['scores'])
            pending=sum(s['approved'] is None and s.get('status')!='按本人填报' for s in person['scores'])
            sheet.cell(total_row,6,f'{unknown}项分数待定；{pending}项材料待核验')
            sheet.cell(total_row,7,'建议分不代表资格已通过或材料已核验')
            names={stage:sorted({d['reviewer'] for k,d in report['decisions'].items() if k.startswith(person['pid']+'|'+person['job']+'|') and d.get('stage')==stage}) for stage in ('初审','复审')}
            sheet.cell(total_row+1,1,'初审记录人：'+ '、'.join(names['初审'])+'    复审记录人：'+'、'.join(names['复审']))
            sheet.print_area=f'A1:G{total_row+1}'
            save_cached(personal,personal_path,{f'E{total_row}':total})
            summary_sheet.cell(n+1,5,total)
            summary_cache[f'E{n+1}']=total
            person['personal_file']=str(personal_path.resolve())
        save_cached(book,out/filename,summary_cache)
