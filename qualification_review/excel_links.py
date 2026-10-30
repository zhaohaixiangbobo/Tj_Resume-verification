"""个人合计与岗位汇总的标准Excel外部链接，保存初始缓存供预览读取。"""
from pathlib import Path
from urllib.parse import quote
from zipfile import ZipFile,ZIP_DEFLATED
import xml.etree.ElementTree as ET
from openpyxl.workbook.external_link.external import (
    ExternalLink,ExternalBook,ExternalSheetNames,ExternalSheetDataSet,ExternalSheetData,ExternalRow,ExternalCell)
from openpyxl.packaging.relationship import Relationship
from openpyxl.workbook.properties import CalcProperties


def link_person(book,summary_sheet,row,relative_path,total_row,value):
    cell=f'E{total_row}'
    data=ExternalSheetData(sheetId=0,row=[ExternalRow(r=total_row,cell=[ExternalCell(r=cell,t='n',v=str(value))])])
    external=ExternalBook(sheetNames=ExternalSheetNames(sheetName=['个人赋分表']),
                          sheetDataSet=ExternalSheetDataSet(sheetData=[data]),id='rId1')
    link=ExternalLink(externalBook=external)
    link.file_link=Relationship(Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/externalLinkPath',
                                Target=quote(Path(relative_path).as_posix()),TargetMode='External',Id='rId1')
    book._external_links.append(link)
    summary_sheet.cell(row,5,f"'[{len(book._external_links)}]个人赋分表'!$E${total_row}")
    summary_sheet.cell(row,5).value='='+summary_sheet.cell(row,5).value
    book.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True,forceFullCalc=True,calcMode='auto')


def save_cached(book,path,cache):
    """cache为单sheet输出的公式单元格初值，不替换公式；Excel后续重新计算。"""
    path=Path(path);book.save(path)
    temporary=path.with_suffix('.cached.xlsx')
    ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    with ZipFile(path) as src,ZipFile(temporary,'w',ZIP_DEFLATED) as dst:
        for item in src.infolist():
            content=src.read(item.filename)
            if item.filename=='xl/worksheets/sheet1.xml':
                root=ET.fromstring(content)
                for cell in root.findall('.//s:c',ns):
                    if cell.attrib.get('r') in cache and cell.find('s:f',ns) is not None:
                        value=cell.find('s:v',ns)
                        if value is None:value=ET.SubElement(cell,'{'+ns['s']+'}v')
                        value.text=str(cache[cell.attrib['r']])
                content=ET.tostring(root,encoding='utf-8',xml_declaration=True)
            dst.writestr(item,content)
    temporary.replace(path)
