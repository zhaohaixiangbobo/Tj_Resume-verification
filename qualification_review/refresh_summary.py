"""读取个人得分列刷新岗位汇总；仅改汇总分数，备份原文件，不运行OCR。"""
from pathlib import Path
from datetime import datetime
from zipfile import ZipFile
import math
import shutil
import xml.etree.ElementTree as ET
import openpyxl

NS='http://schemas.openxmlformats.org/spreadsheetml/2006/main'


def refresh_summary(folder):
    root=Path(folder).resolve()
    results=[]
    backup=root/'内部记录'/('汇总刷新备份_'+datetime.now().strftime('%Y%m%d_%H%M%S_%f'))
    for path in sorted(root.glob('*.xlsx')):
        if path.name.startswith('~$'):continue
        book=openpyxl.load_workbook(path,read_only=True,keep_links=False)
        try:
            if book.sheetnames!=['岗位汇总']:continue
            sheet=book['岗位汇总']
            headers={c.value:c.column for c in sheet[1]}
            if headers.get('身份证号')!=1 or headers.get('建议分小计')!=5:
                raise ValueError(f'{path.name}：汇总表列结构不匹配')
            updates={}
            for row in sheet.iter_rows(min_row=2):
                if row[3].value not in ('符合','待复核'):continue
                ident=str(row[0].value).strip()
                files=list((root/'个人赋分表'/path.stem).glob(ident+'_*.xlsx'))
                if len(files)!=1:raise ValueError(f'{ident}：未找到唯一的个人赋分文件')
                personal=openpyxl.load_workbook(files[0],read_only=True,data_only=False)
                try:
                    ps=personal['个人赋分表']
                    totals=[r[0].row for r in ps if r[0].value=='建议分小计']
                    if len(totals)!=1:raise ValueError(f'{files[0].name}：合计行不唯一')
                    values=[ps.cell(r,5).value for r in range(3,totals[0])]
                    # 空白继续按待定项处理；公式或文字分数必须明确处理，避免读旧缓存。
                    if any(v is not None and (isinstance(v,bool) or not isinstance(v,(int,float)) or not math.isfinite(v)) for v in values):
                        raise ValueError(f'{files[0].name}：建议得分列请填写数值或留空，暂不支持分项公式')
                    updates[f'E{row[0].row}']=sum(v for v in values if v is not None)
                finally:personal.close()
        finally:book.close()
        if not updates:continue
        backup.mkdir(parents=True,exist_ok=True)
        shutil.copy2(path,backup/path.name)
        temporary=path.with_suffix('.refresh.tmp')
        try:
            # 直接修改对应XML单元格，保留其他列、格式、超链接及用户手改内容。
            with ZipFile(path) as src,ZipFile(temporary,'w') as dst:
                for item in src.infolist():
                    content=src.read(item.filename)
                    if item.filename=='xl/worksheets/sheet1.xml':
                        tree=ET.fromstring(content)
                        for cell in tree.iter('{'+NS+'}c'):
                            ref=cell.get('r')
                            if ref in updates:
                                for child in list(cell):cell.remove(child)
                                cell.set('t','n')
                                ET.SubElement(cell,'{'+NS+'}v').text=str(updates[ref])
                        content=ET.tostring(tree,encoding='utf-8',xml_declaration=True)
                    dst.writestr(item,content)
            temporary.replace(path)
        finally:
            if temporary.exists():temporary.unlink()
        results.append({'file':path.name,'scores':updates})
    if not results:raise ValueError('所选目录没有可刷新的岗位汇总，请选择具体批次目录')
    return results


if __name__=='__main__':
    import argparse
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('folder');args=parser.parse_args()
    print(refresh_summary(args.folder))
