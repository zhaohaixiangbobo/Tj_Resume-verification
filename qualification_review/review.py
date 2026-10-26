"""资格审查与赋分试算。读取填报表及岗位表，输出逐项依据、待复核项和业务规则。

默认使用本目录config.json和院校字典。运行过程全部在本机进行。
复杂自由文本采用明确证据计分；信息不足的项目保留待复核，可通过复核表回填。
"""
from __future__ import annotations
import argparse
import calendar
import csv
import hashlib
import json
import re
from datetime import date, datetime
from pathlib import Path
from collections import Counter
import openpyxl
import xlrd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from school_lookup import ROOT, SchoolDictionary, normalize, qs_band
from major_matching import compare_major, parse_allowed, explain, RULE_VERSION

PASS, FAIL, WAIT, NA = '符合', '不符合', '待复核', '不适用'
SCORE_NAMES = ['应聘院校','学历层次','学习形式','政治面貌','毕业年限','在校成绩',
               '奖学金','表彰奖励','计算机等级','英语水平','其他证书','本科学校',
               '人岗匹配','学生工作','科研或实习']
MAX_MASTER = [40,10,10,5,10,2,5,5,1.5,.5,1,4,3,2,1]
MAX_BACHELOR = [40,10,10,5,10,2,5,5,1.5,.5,1,0,5,3,2]
STANDARD_ROWS = [3,13,14,15,16,19,20,26,29,31,32,33,43,45,47]

def text(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()

def interval(value):
    """年月按整月区间表达，避免伪造具体日期。"""
    if isinstance(value, (datetime,date)):
        return value.date() if isinstance(value,datetime) else value, value.date() if isinstance(value,datetime) else value
    m = re.search(r'(20\d{2}|19\d{2})[年./-](\d{1,2})(?:[月./-](\d{1,2}))?', text(value))
    if not m:
        return None
    try:
        y,mo=int(m[1]),int(m[2]); d=int(m[3]) if m[3] else None
        return (date(y,mo,d),date(y,mo,d)) if d else (date(y,mo,1),date(y,mo,calendar.monthrange(y,mo)[1]))
    except ValueError:
        return None

def within(value, start, end):
    span=interval(value)
    if span is None:
        return WAIT
    lo,hi=span
    if hi<start or lo>end:
        return FAIL
    return PASS if start<=lo and hi<=end else WAIT

def degree(value):
    return {'大学本科':'本科','本科':'本科','硕士':'硕士','硕士研究生':'硕士',
            '博士':'博士','博士研究生':'博士'}.get(text(value))

def degree_status(actual, required):
    if actual=='博士' or (actual and required and actual!=required):return FAIL
    return PASS if actual in {'本科','硕士'} and actual==required else WAIT

def major_matches(actual, allowed):
    """兼容调用入口，仅明确符合时返回True；完整结果使用compare_major。"""
    return compare_major(actual, allowed).status == PASS

def education(raw):
    result=[]
    for line in text(raw).replace('\\n','\n').splitlines():
        m=re.search(r'^(.*?)\s+(\d{4}-\d{2})\s+(\d{4}-\d{2})\s+(大学本科|本科|硕士研究生|硕士|博士研究生|博士)\s+(.*?)\s+(?:本科)?(全日制|非全日制)\s*(.*)$',line.strip())
        if m:
            result.append({'school':m[1].strip(),'start':interval(m[2])[0], 'end':interval(m[3])[1],
                           'degree':degree(m[4]),'major':m[5].strip(),'role':m[7].strip(),'raw':line.strip()})
    return result

def party(raw):
    return text(raw) in {'中共党员','中共预备党员','预备党员','中共党员（含预备党员）','中共党员(含预备党员)'}

def awards(raw, kind, periods, school_score, master, cutoff, overseas=False):
    if overseas:
        return 0,False,'境外奖学金和表彰按标准不赋分'
    if text(raw) in {'','无','否'}:
        return 0,False,'未填报该类项目'
    best=0.; pending=False; reasons=[]
    for item in re.split(r'[\n；;]+',text(raw)):
        if not item.strip():
            continue
        item=item.strip(); span=interval(item)
        if re.search(r'未获|未取得|拟申请|计划申请|尚未|预计',item):
            pending=True;reasons.append('取得状态待核对：'+item);continue
        if not span or len(periods)!=1:
            pending=True; reasons.append('待核对日期/最高学历阶段：'+item); continue
        p=periods[0]
        if span[0]>cutoff or span[1]<p['start'] or span[0]>p['end']:
            reasons.append('日期在统计范围外：'+item); continue
        if span[1]>cutoff or span[0]<p['start'] or span[1]>p['end']:
            pending=True; reasons.append('日期跨统计边界：'+item); continue
        score=None
        if kind=='scholarship':
            if '国家励志奖学金' in item:
                score=4 if not master else None
            elif '国家奖学金' in item:
                score=5
            elif '奖学金' not in item:
                pending=True; reasons.append('奖项性质待核对：'+item); continue
            elif '院系级' in item or '院级' in item:
                score=None
            elif '省' in item and ('级' in item or '政府' in item):
                score=None  # 颁发机关与省级认定需人工核对，不能仅靠“省”字。
            elif re.search(r'校级|校[一二三]等',item) or normalize(p['school']) in normalize(item):
                score=2 if school_score==40 else 1 if school_score==35 else .5 if school_score is not None else None
        else:
            exact=re.search(r'优秀毕业生|三好学生|优秀学生干部|优秀共产党员|优秀(?:共青)?团员|优秀(?:共青)?团干部',item)
            if exact:
                score=5 if '国家级' in item else 3 if '省部级' in item else 1 if re.search(r'校级|院级|系部级|院系级',item) else None
            elif re.search(r'竞赛|比赛|大赛|论坛|志愿|社会实践',item):
                score=0
            else:
                score=None
        if score is None:
            pending=True;reasons.append('级别/等同范围待核对：'+item)
        else:
            best=max(best,score);reasons.append(f'候选{score:g}分：'+item)
    # 已明确取得最高5分时，其他待核对奖项不再影响本项得分。
    return best, pending and best<5, '\n'.join(reasons)+'\n取最高一项，不累计。'

def student_score(periods, master, cutoff):
    best=0; pending=False; notes=[]
    for p in periods:
        if p['degree'] not in ({'本科','硕士'} if master else {'本科'}) or p['start']>cutoff:
            continue
        role=p['role']; score=None
        if re.search(r'尚未|未担任|拟任|计划担任',role) and role!='未担任':
            pending=True;notes.append('职务取得状态待核对：'+p['raw']);continue
        if role in {'','无','未担任'}:
            score=0
        elif re.search(r'(?<!副)班长|团支书|团支部书记|(?:学生)?党支部(?:副)?书记',role):
            score=2 if master else 3
        elif re.search(r'(?:学生会|研究生会).*(?:副部长|部长|副主席|主席)',role):
            score=2 if master else 3
        elif re.search(r'(?:学生会|研究生会).*干事|支部委员|班委|副班长|学习委员|生活委员|宣传委员|组织委员|纪律委员|学委',role):
            score=1
        if score is None:
            pending=True;notes.append('职务等同关系待复核：'+p['raw'])
        else:
            best=max(best,score);notes.append(f'候选{score}分：'+p['raw'])
    return best,pending,'\n'.join(notes)

def internship(raw, periods, cutoff):
    """只认定同一条目明确写出实习及日期的本科经历。"""
    if text(raw) in {'无','否'}:
        return 0,False,'明确填报无综合补充经历'
    under=[p for p in periods if p['degree']=='本科']
    if len(under)!=1:
        return None,True,'本科在读时间未能唯一解析'
    p=under[0];notes=[]
    for line in text(raw).splitlines():
        if not re.search(r'实习|社会实践',line) or re.search(r'计划|拟于|尚未|未参加',line):continue
        dates=re.findall(r'(?:19|20)\d{2}[年./-]\d{1,2}(?:[月./-]\d{1,2})?',line)
        if dates:
            spans=[interval(x) for x in dates]
            if any(s is None for s in spans):continue
            lo,hi=spans[0][0],spans[-1][1]
        else:
            years=re.findall(r'((?:19|20)\d{2})年',line)
            if not years:continue
            lo,hi=date(int(years[0]),1,1),date(int(years[-1]),12,31)
        if lo<=hi and p['start']<=lo and hi<=min(p['end'],cutoff):
            return 2,False,'明确本科在读期间的实习：'+line
        notes.append('不在完整可认定期间或跨界：'+line)
    return None,True,'\n'.join(notes) or '需补充本科阶段实习实践起止日期'

def add_sheet(book,name,headers,rows,widths=None):
    s=book.create_sheet(name);s.append(headers)
    for row in rows:
        s.append(list(row))
        # 将来源字符串明确写成文本，避免把QS的“=2”等当Excel公式。
        for c in s[s.max_row]:
            if isinstance(c.value,str): c.data_type='s'
    for c in s[1]:
        c.font=Font(name='微软雅黑',bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='24476A')
        c.alignment=Alignment(wrap_text=True,vertical='center')
    for row in s.iter_rows(min_row=2):
        for c in row:
            c.font=Font(name='微软雅黑',size=10)
            c.alignment=Alignment(vertical='top',wrap_text=True)
            if c.value==WAIT:c.fill=PatternFill('solid',fgColor='FFF2CC')
    for i,h in enumerate(headers,1):
        s.column_dimensions[get_column_letter(i)].width=(widths or {}).get(i,min(42,max(16,len(h)*2+4)))
    s.freeze_panes='A2';s.auto_filter.ref=s.dimensions;s.row_dimensions[1].height=32
    return s

def export_dictionary(schools,out):
    b=openpyxl.Workbook();b.remove(b.active)
    headers=list(schools.rows[0]); add_sheet(b,'国内本科及以上',headers,[[r[h] for h in headers] for r in schools.rows if r['办学层次']!='专科'],{1:32,11:48,12:50,15:50})
    qs=[]
    for r in schools.qs:
        sc=qs_band(r['rank']); rank=int(r['rank'].lstrip('='))
        qs.append([2027,r['rank'],rank,r['chinese'],r['english'],r['location'],
                   '1–50' if sc==40 else '51–100' if sc==35 else '101–150' if sc==30 else '151–200',
                   sc, (sc-20)/5 if sc is not None else None,r['translation_status'],r['source']])
    add_sheet(b,'QS2027前200中英文',['排名版本','原始名次','名次数字','中文名称','英文名称','城市及国家地区','计分档次','应聘院校分','本科院校附加分','译名说明','来源'],qs,{4:38,5:65,6:35,10:48,11:50})
    add_sheet(b,'专科名录备查',headers,[[r[h] for h in headers] for r in schools.rows if r['办学层次']=='专科'])
    notes=[
      ['数据日期','教育部普通高校名单截至2026-06-17；双一流采用第二轮2022名单；QS2027于2026-06-18发布，2026-10-26核对。'],
      ['层次优先级','985 > 211 > 其他双一流 > 其他公办高校；各独立标记仍全部保留。专科不参与本科及以上院校计分。'],
      ['211统计口径','教育部历史原表112个名称；地矿油分别映射两地现行学校后为115个学校实体（含985）；不以统计数量替代逐校依据。'],
      ['学校性质','教育部备注“民办”单列；备注为空的普通学校按公办名录口径整理；合作办学、境外办学特殊条目单列待核验。'],
      ['科研院所及党校','本字典为高校名单，未穷举科研院所和党校。标准中对应按211/普通公办赋分的机构需补充官方身份依据后录入复核表。'],
      ['QS名称','英文名称和排名来自QS官网。中文为整理译名，非QS官方中文版；同名学校结合国家地区及英文名核对。'],
      ['QS范围','仅前200名公开页面的200条记录。未查到学校时待核验，不能由缺席本表推定具体排名。'],
      ['QS源页面检查','源页面第97–100名附近并列编号、115之后跳至118存在序号衔接异常。本表忠实保留原始名次；这些异常均在同一赋分档内，仍建议业务复核。'],
      ['QS发布来源','https://www.qs.com/insights/qs-world-university-rankings-2027-results-table-excel'],
      ['地矿油两地办学依据','https://www.moe.gov.cn/moe_879/moe_1252/s8447/201412/t20141204_180158.html'],
      ['上海体育大学更名依据','https://fzgh.sus.edu.cn/info/1065/1021.htm'],
      ['军医大学更名依据','https://www.mod.gov.cn/gfbw/gfjy_index/4806474.html'],
      ['国防科技大学更名依据','https://www.nudt.edu.cn/xxgk/kdjj/index.htm'],
      ['学校匹配','采用完整校名精确匹配和四条已核实更名映射。独立学院、分校不按母校名称包含关系赋分。']]
    add_sheet(b,'使用说明',['事项','说明或依据'],notes,{1:24,2:110}); b.save(out/'院校层次与QS2027前200中英文.xlsx')

def load_table(path,sheet=None):
    b=openpyxl.load_workbook(path,data_only=True);s=b[sheet] if sheet else b.worksheets[0]
    headers=[text(c.value) for c in s[1]]
    if len(headers)!=len(set(headers)):
        raise ValueError(f'{path.name}存在重复表头')
    return s.title,[(i,{h:v for h,v in zip(headers,row)}) for i,row in enumerate(s.iter_rows(min_row=2,values_only=True),2) if any(v is not None for v in row)],{h:get_column_letter(i+1) for i,h in enumerate(headers)}

def load_reviews(path):
    if path is None:return {}
    _,rows,_=load_table(path,'复核录入')
    result={}
    for _,r in rows:
        if not text(r.get('复核结论')) and r.get('复核得分') is None:continue
        if not text(r.get('复核依据')) or not text(r.get('复核人')):
            raise ValueError('每条复核决定必须填写复核依据和复核人')
        key=(text(r['人员编号']),text(r['岗位编码']),text(r['规则编号']))
        if key in result:raise ValueError(f'重复复核键：{key}')
        result[key]=r
    return result

def run(config_path,person_path=None,review_path=None,out_arg=None):
    cfg=json.loads(config_path.read_text(encoding='utf-8')); base=config_path.parent.parent
    cutoff=date.fromisoformat(cfg['cutoff']); cohort=cfg['cohort']
    start=date.fromisoformat(cfg['graduation_start']);end=date.fromisoformat(cfg['graduation_end'])
    ppath=person_path or base/cfg['personnel_file'];jpath=base/cfg['jobs_file']
    out=out_arg or ROOT/'outputs'/datetime.now().strftime('%Y%m%d_%H%M%S_%f');out.mkdir(parents=True,exist_ok=False)
    print('[1/4] 读取输入与院校字典',flush=True)
    schools=SchoolDictionary();psheet,persons,pcols=load_table(ppath,cfg['personnel_sheet'])
    jsheet,jobs,_=load_table(jpath)
    required=['姓名','身份证号码','岗位编码','出生日期','考生类别','政治面貌','最高学历','最高学位','最高学历毕业院校','最高学历毕业时间','最高学历所学专业','最高学历学习形式','学习经历','奖学金','表彰奖励','英语等级','计算机等级']
    for h in required:
        if h not in pcols:raise ValueError('人员表缺少列：'+h)
    jobmap={}
    for i,j in jobs:
        key=text(j['岗位编码'])
        if not key or key in jobmap:raise ValueError('岗位编码空白或重复：'+key)
        jobmap[key]=(i,j)
    reviews=load_reviews(review_path);used=set();details=[];summaries=[];forms=[];major_details=[]
    ids=Counter(text(p['身份证号码']) for _,p in persons)
    std=[]
    for file in ['管理岗标准.xls','业务岗标准.xls']:
        path=base/file;digest=hashlib.sha256(path.read_bytes()).hexdigest()
        if digest!=cfg['standard_sha256'][file]:raise ValueError(file+'已变化，请先审核新标准并更新程序规则及配置中的哈希')
        sh=xlrd.open_workbook(path).sheet_by_index(0)
        for i in range(sh.nrows):std.append([file,sh.name,i+1,*[text(v) for v in sh.row_values(i)],'',''])
    manifest={'config':cfg,'major_rule_version':RULE_VERSION,
              'major_engine_sha256':hashlib.sha256((ROOT/'major_matching.py').read_bytes()).hexdigest(),
              'personnel_sha256':hashlib.sha256(ppath.read_bytes()).hexdigest(),
              'jobs_sha256':hashlib.sha256(jpath.read_bytes()).hexdigest(),
              'domestic_sha256':hashlib.sha256((ROOT/'data/domestic_schools.csv').read_bytes()).hexdigest(),
              'qs_sha256':hashlib.sha256((ROOT/'data/qs_2027_bilingual.json').read_bytes()).hexdigest()}
    for n,(r,p) in enumerate(persons,1):
        name=text(p['姓名']); jobid=text(p['岗位编码']); ident=text(p['身份证号码'])
        pid=hashlib.sha256(ident.encode()).hexdigest()[:16] if ident else f'无证件号-源行{r}'
        fingerprint=hashlib.sha256(json.dumps([p,jobmap.get(jobid),manifest],ensure_ascii=False,default=str,sort_keys=True).encode()).hexdigest()
        rows=[]
        def add(code,label,status,score,maxscore,reason,columns):
            raw='\n'.join(f'{c}：{text(p.get(c))}' for c in columns)
            loc='；'.join(f'{ppath.name}/{psheet}!{pcols[c]}{r}' for c in columns if c in pcols)
            key=(pid,jobid,code)
            if key in reviews:
                a=reviews[key]
                if text(a['输入指纹'])!=fingerprint:raise ValueError(f'{name}/{code}复核对应的输入已变化，请重新复核')
                if code.startswith('Q'):
                    if a['复核结论'] not in {PASS,FAIL,WAIT,NA}:raise ValueError('复核结论只允许符合、不符合、待复核、不适用')
                    status=a['复核结论']
                else:
                    value=a['复核得分']
                    allowed={0,maxscore}
                    if code in {'S01'}:allowed={20,25,30,35,40}
                    if code=='S07':allowed={0,.5,1,2,3,5} | ({4} if not master else set())
                    if code=='S08':allowed={0,1,3,5}
                    if code=='S09':allowed={0,.5,1.5}
                    if code=='S12':allowed={0,1,2,3,4} if master else {0}
                    if code=='S14':allowed={0,1,maxscore}
                    if not isinstance(value,(float,int)) or value not in allowed:raise ValueError(f'{name}/{code}分数不在允许档位')
                    score=float(value);status='已复核'
                reason+=f"\n人工复核（{a['复核人']}）：{a['复核依据']}";used.add(key)
            rows.append([pid,name,jobid,code,label,status,score,maxscore,reason,raw,loc])
            forms.append([pid,name,jobid,code,label,fingerprint,status,score,'',None,'',''])
        d=degree(p['最高学历'])
        js=jobmap.get(jobid);j=js[1] if js else {};jdegree=degree(j.get('学历要求'))
        master=jdegree=='硕士' if jdegree else d=='硕士';limits=MAX_MASTER if master else MAX_BACHELOR
        add('Q00','记录和岗位关联',PASS if js and ident and ids[ident]==1 else WAIT,None,None,
            '按岗位编码关联；重复报名、无证件号或岗位未命中需复核', ['岗位编码','姓名'])
        cat=text(p['考生类别']);overseas='境外' in cat or '国外' in cat or '留学' in cat
        domestic='国内' in cat
        add('Q01','考生类别',PASS if domestic or overseas else WAIT,None,None,'按考生类别分流，国家及学校名称仅辅助核验',['考生类别'])
        grad=interval(p['最高学历毕业时间'])
        add('Q02','毕业日期范围',within(p['最高学历毕业时间'],start,end),None,None,
            f'{start}至{end}（含两端）；暂未毕业按预计日期预审',['最高学历毕业时间','考生类别'])
        add('Q03','普通高校及培养方式',WAIT if domestic else NA,None,None,
            '补充国内普通高校、非委培、非定向声明；全日制字段不能独立证明全部条件',['最高学历学习形式','最高学历毕业院校'])
        add('Q04','未落实工作单位',WAIT,None,None,'补充未落实工作单位声明；实习、档案存放地点不能替代就业状态',['考生类别'])
        pair={'本科':'学士','硕士':'硕士','博士':'博士'}
        add('Q05','学历学位组合',PASS if d and text(p['最高学位'])==pair[d] else WAIT,None,None,'核对最高学历对应的学位名称',['最高学历','最高学位'])
        add('Q06','学历学位取得及认证',WAIT,None,None,
            '补充学位取得日期、教育部门认证完成情况和日期' if overseas else '补充国家承认、已取得或预计在毕业期限内取得学历及学位的声明',['最高学历','最高学位','最高学历毕业时间'])
        birth=interval(p['出生日期']);age18=date(cutoff.year-18,cutoff.month,cutoff.day)
        age_status=within(p['出生日期'],date(1900,1,1),age18)
        iddate=None
        if re.fullmatch(r'\d{17}[\dXx]',ident):
            try:iddate=datetime.strptime(ident[6:14],'%Y%m%d').date()
            except ValueError:pass
        if iddate and birth and not birth[0]<=iddate<=birth[1]:age_status=WAIT
        add('Q07','年满18周岁',age_status,None,None,f'以{cutoff}为截止日，出生日期须不晚于{age18}；出生日期与身份证交叉核对',['出生日期'])
        threshold=date(1999,1,1) if d=='硕士' else date(2002,1,1)
        add('Q08','出生日期下限',within(p['出生日期'],threshold,cutoff) if d in {'本科','硕士'} else NA,None,None,f'{d}出生日期须不早于{threshold}',['出生日期','最高学历'])
        add('Q09','学历与岗位精确对应',degree_status(d,jdegree),None,None,
            f'最高学历{d}；岗位学历{jdegree}。本科报本科，硕士报硕士。岗位来源：{jpath.name}/{jsheet}!F{js[0] if js else "?"}',['最高学历','最高学位','岗位编码'])
        allowed=j.get('研究生专业要求' if master else '本科专业要求')
        match=compare_major(p['最高学历所学专业'],allowed)
        jobloc=f'{jpath.name}/{jsheet}!{"H" if master else "G"}{js[0] if js else "?"}'
        add('Q10','最高学历主专业',match.status,None,None,
            f'允许专业原文：{text(allowed)}\n'+explain(match)+f'\n第二专业和研究方向不替代主专业。岗位来源：{jobloc}',
            ['最高学历所学专业','岗位编码'])
        major_details.append([pid,name,jobid,text(p['最高学历所学专业']),match.actual,
                              '、'.join(match.allowed),'、'.join(match.matched),match.kind,match.status,rows[-1][5],match.reason,jobloc])
        other=text(j.get('其他资格条件要求'))
        add('Q11','岗位党员条件',(PASS if party(p['政治面貌']) else FAIL if text(p['政治面貌']) else WAIT) if '中共党员' in other else NA,None,None,
            '接受中共党员及预备党员；岗位其他条件：'+other,['政治面貌','岗位编码'])
        axkey=next((h for h in p if h.startswith('法律职业资格证书')),None);ax=text(p.get(axkey))
        legal=bool(re.search(r'(?:具有|持有|取得|获得|已获).{0,8}法律职业资格(?:证书|证)|法律职业资格(?:证书|证)[（(]A证|法律职业资格A证',ax)) and not bool(re.search(r'未(?:取得|获得)|尚未|拟申请|待领取',ax))
        add('Q12','岗位法律资格证',(PASS if legal else WAIT) if '法律职业资格' in other else NA,None,None,
            '按明确持证自述预审，考试通过不能独立替代取得证书；岗位其他条件：'+other,[axkey] if axkey else [])
        periods=education(p['学习经历']);high=[e for e in periods if e['degree']==d]
        cross=PASS if len(high)==1 else WAIT
        discrepancies=[]
        if len(high)!=1:discrepancies.append('最高学历经历未能唯一解析')
        if any(e['degree']=='博士' for e in periods) and d!='博士':
            cross=WAIT;discrepancies.append('经历中存在更高的博士学历')
        if len(high)==1:
            if normalize(high[0]['school'])!=normalize(p['最高学历毕业院校']):
                cross=WAIT;discrepancies.append('AG学校与Y学校名称不同')
            pm=compare_major(high[0]['major'],p['最高学历所学专业'])
            if pm.status!=PASS:cross=WAIT
            discrepancies.append(f'AG专业={high[0]["major"]}；AA专业={text(p["最高学历所学专业"])}\n'+explain(pm))
        if grad and len(high)==1 and (grad[0].year,grad[0].month)!=(high[0]['end'].year,high[0]['end'].month):
            cross=WAIT;discrepancies.append('AG与Z毕业年月不同')
        add('Q13','最高学历经历一致性',cross,None,None,'核对学校、专业、毕业年月；疑似修饰词差异提示复核。\n'+'\n'.join(discrepancies),['学习经历','最高学历毕业院校','最高学历所学专业','最高学历毕业时间'])
        leftover=re.sub(r'中共党员|取得法律职业资格证书|[、,，;；。\s]','',other)
        add('Q14','岗位其他附加条件',WAIT if leftover else NA,None,None,
            '岗位表其他资格条件原文：'+(other or '无')+'；未被党员和法律证书规则覆盖的要求逐项人工复核',['岗位编码'])
        ss,snote=schools.lookup(p['最高学历毕业院校'],overseas)
        standard='管理岗标准.xls' if master else '业务岗标准.xls';scorecol='E' if master else 'D'
        def score(i,value,reason,columns,pending=False):
            add(f'S{i:02}',SCORE_NAMES[i-1],WAIT if pending or value is None else '按填报试算',value,limits[i-1],
                reason+f'\n标准：{standard}/天津市局!{scorecol}{STANDARD_ROWS[i-1]}，完整条款见核验规则工作簿。',columns)
        score(1,ss,snote,['最高学历毕业院校','考生类别'])
        score(2,10 if d in {'本科','硕士'} and d==jdegree else None,'按最高学历与岗位对应标准计分',['最高学历','岗位编码'])
        score(3,10 if p['最高学历学习形式']=='全日制' else 0 if p['最高学历学习形式']=='非全日制' else None,'全日制10分',['最高学历学习形式'])
        score(4,5 if party(p['政治面貌']) else 0 if text(p['政治面貌']) else None,'党员含预备党员5分',['政治面貌'])
        score(5,10 if grad and grad[0].year==cohort else 5 if grad and start.year<=grad[0].year<cohort else None,f'招聘届别固定为{cohort}；当年10分，往年5分，身份另核验',['最高学历毕业时间'])
        gradecol=next(h for h in p if h.startswith('在读期间有无'))
        score(6,2 if p[gradecol]=='否' else None,'无补考重修不合格记录2分；有记录者核对学校正式审批例外',[gradecol])
        for i,kind,field,flag in [(7,'scholarship','奖学金','高等教育期间是否取得奖学金'),(8,'honor','表彰奖励','高等教育期间是否获得表彰奖励')]:
            value,pend,why=awards(p[field],kind,high,ss,master,cutoff,overseas)
            if not text(p[field]) and text(p.get(flag))!='否':pend=True;why+='；明细为空，需核对'
            score(i,value,why,[field,'学习经历'],pend)
        comp=text(p['计算机等级']);cs=1.5 if re.search('三级|四级',comp) else .5 if '二级' in comp else 0 if comp in {'无','未取得'} else None
        score(9,cs,'按填报全国计算机等级；“其他”需补充具体级别。自述取得时间按导出日口径',['计算机等级'])
        eng=text(p['英语等级']);es=.5 if '六级' in eng and '425' in eng else None
        if es is None:
            for pattern,threshold in [(r'雅思\D{0,5}(\d+(?:\.\d+)?)',6.5),(r'托福\D{0,5}(\d+(?:\.\d+)?)',90)]:
                m=re.search(pattern,ax)
                if m and float(m[1])>=threshold and not re.search(r'未|目标|计划|拟',m[0]):es=.5
        score(10,es,'六级425、专四/专八60、雅思6.5、托福90任一达标。未明确达标的补充核验',['英语等级']+([axkey] if axkey else []))
        cpa=bool(re.search(r'(?:取得|获得|持有).{0,8}注册会计师全科合格证',ax)) and not bool(re.search(r'未取得|未获得|尚未|拟申请',ax))
        score(11,1 if legal or cpa else None,'取得法职资格证或CPA全科合格证1分；未明确填报证书情况待复核',[axkey] if axkey else [])
        under=[e for e in periods if e['degree']=='本科'];us=None;unote='本科经历未能唯一解析'
        if len(under)==1:
            raw,unote=schools.lookup(under[0]['school']);us=(raw-20)/5 if raw is not None else None
        score(12,us if master else 0,unote if master else '业务岗无本科学校附加分',['学习经历'])
        score(13,None,'按标准由业务部门综合评定并填写依据；管理岗0或3分，业务岗0或5分',[axkey] if axkey else [])
        st,sp,sr=student_score(periods,master,cutoff)
        # AX可能记载比AG更高职务，保留复核，不能因AG未提及便漏计。
        if st<limits[13] and re.search(r'班长|团支|学生会|研究生会|党支部|副部长|学生工作',ax):sp=True;sr+='\nAX可能另含学生工作，需核对在读阶段及职务。'
        score(14,st,sr or '学习经历未解析到可认定学生职务',['学习经历']+([axkey] if axkey else []),sp or not periods)
        if master:
            research_value=0 if ax in {'无','否'} else None
            score(15,research_value,'硕士期间科研需核对作者/发明人顺位、核心级别、发表或授权日期、竞赛级别',[axkey] if axkey else [])
        else:
            iv,ip,ir=internship(ax,periods,cutoff)
            score(15,iv,ir,[axkey,'学习经历'] if axkey else ['学习经历'],ip)
        q=[a for a in rows if a[3].startswith('Q')];scores=[a for a in rows if a[3].startswith('S')]
        conclusion=FAIL if any(a[5]==FAIL for a in q) else WAIT if any(a[5]==WAIT for a in q) else PASS
        subtotal=sum(a[6] or 0 for a in scores);pending=sum(a[5]==WAIT for a in scores)
        assert 0<=subtotal<=100
        summaries.append([pid,name,jobid,text(p.get('岗位名称')),conclusion,'管理岗' if master else '业务岗',subtotal,
                          pending,subtotal if pending==0 else None,'全部项目已试算' if pending==0 else '部分项目待复核，小计不作最终排名',r,
                          '\n'.join(a[4]+':'+a[5] for a in q if a[5] in {FAIL,WAIT})])
        details.extend(rows)
        print(f'  [{n}/{len(persons)}] {n/len(persons):.0%} 已完成逐项核验',flush=True)
    unused=set(reviews)-used
    if unused:raise ValueError(f'{len(unused)}条复核未对应本次人员或规则，已停止输出')
    print('[3/4] 导出结果、规则、院校字典',flush=True)
    b=openpyxl.Workbook();b.remove(b.active)
    add_sheet(b,'人员汇总',['人员编号','姓名','岗位编码','岗位名称','资格预审','采用标准','已识别试算小计','待复核赋分项数','完整试算分','分数状态','人员源行','资格问题'],summaries,{10:44,12:75})
    add_sheet(b,'专业匹配明细',['人员编号','姓名','岗位编码','AA填报原文','规范化专业','岗位专业清单','命中或候选专业','匹配类型','程序判断','复核后结论','原因','岗位来源'],major_details,{4:32,5:32,6:90,7:48,11:80,12:65})
    add_sheet(b,'逐项依据',['人员编号','姓名','岗位编码','规则编号','项目','状态','已识别分值','项目上限','处理依据','原始填报','来源位置'],details,{9:80,10:95,11:65})
    fs=add_sheet(b,'复核录入',['人员编号','姓名','岗位编码','规则编号','项目','输入指纹','当前状态','当前分数','复核结论','复核得分','复核依据','复核人'],forms,{5:26,6:26,11:80})
    dv=DataValidation(type='list',formula1='"符合,不符合,待复核,不适用"');fs.add_data_validation(dv);dv.add(f'I2:I{fs.max_row}')
    add_sheet(b,'运行说明',['事项','说明'],[
        ['截止日',str(cutoff)],['输入快照日期',cfg['snapshot_date']],['试算依据','按自述填报信息，材料真实性未核验。奖项逐条核对日期及学历阶段。'],
        ['分数解释','已识别小计包含已识别证据；待复核项目可能已有部分得分。完整试算分仅在各评分项均明确时显示。资格不符合人员仍列出诊断性试算，不列入合格排名。'],
        ['复核操作','在复核录入表填写I/J/K/L列，另存工作簿后用 --review 指定。资格填结论；赋分填允许档位分数；每项必须有依据和复核人。输入变化须重做复核。'],
        ['复核时间要求','对证书、职务、奖项和实习，复核依据写明取得/发生日期、所处学历阶段及适用条款。'],
        ['重复报名','同证件号记录进入待复核；请先处理重复报名，复核表键为人员编号+岗位编码+规则编号。'],
        ['人工判断','人岗匹配和复杂科研/实习需业务部门判断；政治面貌仅使用本人填报，性别民族身高体重不参与赋分。']],{2:110})
    b.save(out/'资格审查与赋分试算.xlsx')
    export_rules(cfg,base,out)
    export_dictionary(schools,out)
    (out/'运行记录.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    with (out/'人员汇总.csv').open('w',encoding='utf-8-sig',newline='') as f:
        writer=csv.writer(f);writer.writerow([c.value for c in b['人员汇总'][1]]);writer.writerows(summaries)
    print(f'[4/4] 完成：{len(persons)}人；资格状态{dict(Counter(r[4] for r in summaries))}\n输出目录：{out}',flush=True)
    return out

def export_rules(cfg,base,out):
    b=openpyxl.Workbook();b.remove(b.active)
    rules=[
        ['Q00','记录关联','F岗位编码、K身份证号码','岗位表D','岗位编码精确关联，重复或缺失待复核'],
        ['Q01','身份分流','U考生类别','国内/境外','按明确类别分流'],
        ['Q02','毕业时间','Z毕业日期、AG经历',cfg['graduation_start']+'至'+cfg['graduation_end'],'含两端；预计毕业可预审；年月用整月区间'],
        ['Q03','培养身份','AE学习形式、Y学校','补充声明','国内普通高校，非委培、非定向，缺声明待复核'],
        ['Q04','就业状态','U、补充声明','未落实工作单位','实习和档案信息不推断就业状态'],
        ['Q05','学历学位','W学历、X学位','本科/学士；硕士/硕士','核对组合'],
        ['Q06','取得与认证','W/X/Z、补充声明','毕业期限内取得学历学位；境外完成教育部门认证','补充日期和声明，缺失待复核'],
        ['Q07','18周岁','L出生日期、K身份证交叉核对',cfg['cutoff'],'截至本日满18周岁'],
        ['Q08','出生下限','L、W','本科2002-01-01；硕士1999-01-01','含边界；博士按Q09排除'],
        ['Q09','学历对应','W/X、岗位表F','本科报本科，硕士报硕士','博士不接受；硕士不能报本科岗'],
        ['Q10','主专业','AA、AG；岗位表G/H','精确命中/疑似待复核/不匹配','本科用G，硕士用H；修饰词、包含或共同前缀待复核；保留法律限定；中英文译名不自动等同，详见专业匹配规则'],
        ['Q11','党员要求','S；岗位表I','含预备党员','按本人明确填报'],
        ['Q12','法律资格','AX；岗位表I','明确取得法律职业资格证','考试通过与持证区分'],
        ['Q13','经历一致性','AG与W/Y/Z/AA','学校、主专业、毕业年月一致','识别失败或差异待复核'],
        ['Q14','其他附加条件','岗位表I','党员和法律资格以外要求','新出现的要求保留人工复核'],
    ]
    add_sheet(b,'资格核验规则',['编号','项目','取值列','标准或参数','处理方式','业务审核意见','审核人'],[r+['',''] for r in rules],{3:34,4:55,5:85,6:50})
    jsheet,jobs,_=load_table(base/cfg['jobs_file'])
    jobrows=[]
    for i,j in jobs:
        raw_g=text(j.get('本科专业要求'));raw_h=text(j.get('研究生专业要求'))
        g=parse_allowed(raw_g);h=parse_allowed(raw_h)
        issue='；'.join(x.issue for x,raw in [(g,raw_g),(h,raw_h)] if raw and x.issue) or '已解析；空白专业列不适用'
        jobrows.append([text(j['岗位编码']),text(j.get('用人单位')),text(j.get('岗位名称')),text(j.get('学历要求')),
                        raw_g,'、'.join(g.items),raw_h,'、'.join(h.items),issue,f'{cfg["jobs_file"]}/{jsheet}!G{i}:H{i}'])
    add_sheet(b,'各岗位专业清单',['岗位编码','用人单位','岗位名称','学历要求','本科专业原文','本科规范化清单','研究生专业原文','研究生规范化清单','解析说明','来源'],jobrows,{2:42,5:75,6:75,7:85,8:85,9:42,10:60})
    examples=[('网络工程专业','网络工程'),('会计专业硕士','会计学、会计'),('会计学','会计'),
              ('网络工程','网络工程'),('马克思主义理论','马克思主义哲学'),
              ('法学（030101K）','法学'),('计算机科学与技术（软件方向）','计算机科学与技术'),
              ('法律（非法学）','法律（法学）'),('法律(非法学)','法律（法学）、法律（非法学）'),
              ('法律','法律（法学）'),('法律（非法学）','法律'),
              ('Accounting','会计'),('Computer Science','computer science'),
              ('计算机科学与技术','软件工程,计算机科学\\n与技术,人工智能'),
              ('软件工程','软件工程\n会计'),('会计','财务管理')]
    erows=[]
    for actual,allow in examples:
        m=compare_major(actual,allow);erows.append([actual,allow,m.status,m.kind,'、'.join(m.matched),m.reason,'',''])
    add_sheet(b,'专业匹配规则',['填报示例','岗位要求示例','程序结果','规则类型','命中或候选','处理说明','业务审核意见','审核人'],erows,{1:35,2:55,5:48,6:100,7:50})
    processing=[
        '完整校名查教育部字典；境外查QS2027前200；未知待复核',
        '学历与岗位精确对应，符合本标准学历10分', '全日制10分', '中共党员含预备5分',
        '2027年毕业10分，2025/2026年5分，身份单独核验',
        'AF否=2分；AF是先核对学校正式审批例外',
        'AJ逐条日期对最高学历在读期及截止日；最高一次；励志仅业务4分；院系级和颁发机关歧义复核',
        'AM仅最高学历期间列举荣誉，最高一次；竞赛不计本荣誉项；优秀研究生等同关系复核',
        'AO全国计算机三级以上1.5、二级0.5；其他补充',
        'AN及AX任一英语阈值达标0.5；未明确部分需人工复核',
        'AX明确持证法律或CPA全科1分；未明确补充',
        '管理岗AG本科院校4/3/2/1/0；业务岗不加此项',
        '业务人工评价，管理0/3，业务0/5；必须写依据',
        'AG和AX；管理本科及硕士阶段0/1/2，业务本科阶段0/1/3；未列职务等同关系复核',
        '管理：仅硕士科研0/1，核对作者位次和公开发表/授权等；业务：仅本科实习实践0/2；人工记录依据',
    ]
    mapping=['Y/U/AG','W/X/F','AE/AG','S','Z/U','AF','AH/AI/AJ/AG/Y','AK/AL/AM/AG','AO/AX','AN/AX','AX','AG','AX及人工','AG/AX','AX/AG/AM']
    add_sheet(b,'赋分处理规则',['编号','项目','人员列','管理上限','业务上限','程序处理','业务审核意见','审核人'],
              [[f'S{i+1:02}',SCORE_NAMES[i],mapping[i],MAX_MASTER[i],MAX_BACHELOR[i],processing[i],'',''] for i in range(15)],{6:100,7:50})
    for fn in ['管理岗标准.xls','业务岗标准.xls']:
        s=xlrd.open_workbook(base/fn).sheet_by_index(0)
        data=[]
        for i in range(s.nrows):data.append([i+1,*s.row_values(i)])
        add_sheet(b,fn[:-4]+'原文',['源Excel行']+[get_column_letter(i+1)+'列原文' for i in range(s.ncols)],data,{i:40 for i in range(2,s.ncols+2)})
    add_sheet(b,'参数与业务确认',['事项','采用口径','业务审核意见','审核人'],[
        ['年龄、获奖、取证截止',cfg['cutoff']+'；“10月底”按本次导出所属2026年解释','',''],
        ['填报快照',cfg['snapshot_date']+'；明确自述已持有的证书按截至导出时已取得预审，具体日期可人工补充','',''],
        ['QS年度','采用截至本次核对最新已发布的QS2027主榜；替代原标准“毕业当年最近一次”口径','',''],
        ['试算与正式分','当前按填报文字试算；正式分仍需按标准完成材料核验。待复核分数不伪装成完整总分','',''],
        ['身份声明','普通高校、非委培非定向、未落实工作、学历学位取得和境外认证缺列，逐人复核填写','',''],
        ['管理励志奖学金','管理标准对应行是/，保留待复核；业务标准明确为4分','',''],
        ['业务标准已更换','采用当前新版，实习2分、学生工作3分、人岗匹配5分','',''],
        ['院校特殊情况','合作办学、科研院所、党校和未知/旧校名先核对；未找到不直接按最低档','',''],
        ['人工评分','人岗匹配、复杂科研/实习及未能可靠解析的证据人工判定；每项需依据及复核人','',''],
        ['排名','本版输出试算分及待复核项，不生成录取名单或缺项情况下的排名','',''],
        ['格式核对','按用户要求未进行视觉版式核对，仅执行结构、内容及核心逻辑测试','','']],{2:110,3:50})
    b.save(out/'核验规则_业务审核.xlsx')

if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--config',type=Path,default=ROOT/'config.json')
    parser.add_argument('--input',type=Path,help='另一个同结构人员Excel')
    parser.add_argument('--review',type=Path,help='填好复核录入Sheet的结果Excel')
    parser.add_argument('--output',type=Path,help='全新输出目录，避免覆盖已有结果')
    a=parser.parse_args()
    run(a.config.resolve(),a.input.resolve() if a.input else None,a.review.resolve() if a.review else None,a.output.resolve() if a.output else None)
