"""桌面流程必要边界测试：资格分流、强制凭证、境外日期、复核导入及模板清除。"""
import contextlib
import copy
import io
import json
import sys
import unittest
import uuid
from pathlib import Path
import openpyxl
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from candidate_rules import candidate_category,CATEGORIES
from desktop_workflow import validate_decision,run_batch
from desktop_review import apply_decisions,import_reviews
from desktop_export import template,table,HEADERS,suggested,clean_reason,brief_info
from evidence_pipeline import digest
from review import run,ROOT


class DesktopTests(unittest.TestCase):
    def setUp(self):
        self.folder=ROOT/'tests/_runs'/('desktop_'+uuid.uuid4().hex);self.folder.mkdir(parents=True)
        self.photo=self.folder/'虚构证明.jpg';self.photo.write_bytes(b'synthetic evidence fixture')
        self.person={'pid':'P1','job':'001','fingerprint':'v1','master':False,
            'fields':{'考生类别':CATEGORIES[3]},'material_manifest':[{'path':str(self.photo),'sha256':digest(self.photo)}],
            'base_qualification':{'Q11':'待复核','Q12':'待复核'},'qualification':[{'code':'Q12'}],'scores':[{'code':'S04'}]}
        self.decision={'fingerprint':'v1','reviewer':'初审甲','reason':'已核对虚构材料及本人归属',
            'conclusion':'符合','score':5,'proofs':[str(self.photo)],'proof_checked':True,'stage':'初审',
            'degree_date':'2026-07-01','certification_date':'2027-07-31'}

    def test_four_categories(self):
        for i,c in enumerate(CATEGORIES):
            r=candidate_category(c);self.assertEqual(r['overseas'],i>=2);self.assertEqual(r['domestic'],i<2)
        self.assertFalse(candidate_category('国内？')['domestic'])

    def test_mandatory_proof_and_old_hash(self):
        validate_decision(self.decision,self.person,'Q12')
        for change in [{'proofs':[]},{'proof_checked':False},{'fingerprint':'old'},{'conclusion':'不适用'}]:
            with self.assertRaises(ValueError):validate_decision(dict(self.decision,**change),self.person,'Q12')
        with self.assertRaises(ValueError):validate_decision(dict(self.decision,proofs=[]),self.person,'S04')
        self.photo.write_bytes(b'changed')
        with self.assertRaises(ValueError):validate_decision(self.decision,self.person,'Q12')

    def test_overseas_deadline(self):
        validate_decision(self.decision,self.person,'Q06')
        for change in [{'certification_date':'2027-08-01'},{'degree_date':'2024-12-31'},{'certification_date':''}]:
            with self.assertRaises(ValueError):validate_decision(dict(self.decision,**change),self.person,'Q06')

    def test_review_stages_and_atomic_import(self):
        report={'persons':[self.person]};path=self.folder/'reviews.json'
        apply_decisions(report,[(('P1','001'),'Q12',self.decision)],path)
        with self.assertRaises(ValueError):apply_decisions(report,[(('P1','001'),'Q12',dict(self.decision,stage='复审'))],path)
        apply_decisions(report,[(('P1','001'),'Q12',dict(self.decision,stage='复审',reviewer='复审乙'))],path)
        stored=json.loads(path.read_text(encoding='utf-8'))['decisions']['P1|001|Q12']
        self.assertEqual(len(stored['history']),1)
        apply_decisions(report,[(('P1','001'),'Q12',dict(self.decision,stage='复审',reviewer='复审乙'))],path)
        self.assertEqual(len(json.loads(path.read_text(encoding='utf-8'))['decisions']['P1|001|Q12']['history']),1)
        book=openpyxl.Workbook();book.remove(book.active)
        table(book,'复核录入',HEADERS,[['P1','001','虚构人员','S04','政治面貌','v1','','',None,'',5,'看过证明','甲','', '是','初审','','','']])
        file=self.folder/'invalid.xlsx';book.save(file);before=path.read_bytes()
        with self.assertRaises(ValueError):import_reviews(report,file,path)
        self.assertEqual(path.read_bytes(),before)

    def test_template_example_values_removed(self):
        for master in (False,True):
            b=openpyxl.Workbook();s=template(b,'模板',master)
            self.assertTrue(all(s.cell(r,c).value is None for r in range(3,16 if master else 15) for c in (4,5,6,7)))
            self.assertEqual(s['B3'].value,'院校层次')

    def test_gate_and_three_person_export(self):
        cfg=json.loads((ROOT/'config.json').read_text(encoding='utf-8'))
        source=ROOT.parent/cfg['personnel_file']
        with contextlib.redirect_stdout(io.StringIO()):base=run(ROOT/'config.json',data_only=True)
        bad={r[0] for r in base['summaries'] if r[4]=='不符合'}
        self.assertFalse(any(r[0] in bad and r[3].startswith('S') for r in base['details']))
        chosen=[]
        for status,kind in [('待复核','业务岗'),('待复核','管理岗'),('不符合','业务岗')]:
            summary=next(r for r in base['summaries'] if r[4]==status and r[5]==kind)
            p=next(p for row,p in base['persons'] if row==summary[10]);chosen.append(str(p['身份证号码']))
        options={'input':str(source),'materials':str(source.parent),'selected_ids':chosen,'output':str(self.folder/'batch'),
                 'ocr':False,'llm':False,'max_pages':1}
        with contextlib.redirect_stdout(io.StringIO()):report=run_batch(options)
        self.assertEqual(len(report['persons']),3)
        failed=next(p for p in report['persons'] if p['conclusion']=='不符合')
        self.assertEqual(failed['scores'],[]);self.assertEqual(failed['evidence']['pages'],[])
        out=self.folder/'batch'
        self.assertFalse((out/'批次汇总.xlsx').exists())
        workbooks=list(out.glob('*.xlsx'))
        self.assertEqual(len(workbooks),len({(p['job'],p['unit'],p['job_name'],p['master']) for p in report['persons']}))
        total_people=0
        for path in workbooks:
            b=openpyxl.load_workbook(path,data_only=True)
            self.assertEqual(b.sheetnames[0],'岗位汇总')
            self.assertEqual(len(b.sheetnames),1)
            self.assertEqual(b['岗位汇总']['A1'].value,'身份证号')
            self.assertIsInstance(b['岗位汇总']['A2'].value,str)
            self.assertEqual(len(b['岗位汇总']['A2'].value),18)
            total_people+=b['岗位汇总'].max_row-1
            for s in b:
                self.assertFalse(any(isinstance(c.value,str) and ('标准：' in c.value or c.value in {'人员编号','岗位赋分文件'}) for row in s for c in row))
            for s in list(b)[1:]:
                self.assertEqual(s['E2'].value,'建议得分')
                self.assertIsNotNone(s['E3'].value)
        self.assertEqual(total_people,3)
        personal=list((out/'个人赋分表').rglob('*.xlsx'))
        self.assertEqual(len(personal),2)
        for path in personal:
            self.assertEqual(len(path.stem.split('_')[0]),18)
            b=openpyxl.load_workbook(path,data_only=True)
            self.assertEqual(b.sheetnames,['个人赋分表'])
            self.assertIsNotNone(b['个人赋分表']['E3'].value)
            self.assertFalse(any('材料信息一致（建议）' in str(c.value) for row in b.active for c in row))
        self.assertTrue((out/'内部记录/人工复核录入.xlsx').exists())
        self.assertEqual(len(report['stale_reviews']),0)

    def test_suggested_scores_and_short_information(self):
        self.assertEqual(suggested({'tentative':25,'approved':None}),25)
        self.assertEqual(suggested({'tentative':25,'approved':0}),0)
        self.assertIsNone(suggested({'tentative':None,'approved':None}))
        self.assertEqual(clean_reason('按条件计分\n标准：管理岗标准.xls/天津市局!E26，完整条款见核验规则工作簿。'),'按条件计分')
        person={'fields':{'计算机等级':'计算机二级','英语等级':'六级','法律职业资格证书，实习实践':'公司实习经历'},'scores':[]}
        info=brief_info(person,['S09','S10','S11'])
        self.assertIn('计算机二级',info);self.assertIn('六级',info);self.assertNotIn('实习',info)


if __name__=='__main__':unittest.main()
