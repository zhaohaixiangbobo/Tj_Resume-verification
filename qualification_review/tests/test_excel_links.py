"""检查跨工作簿引用、个人SUM公式和初始缓存，避免预览显示空白分数。"""
import sys
import unittest
import uuid
from pathlib import Path
from zipfile import ZipFile
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import openpyxl
from excel_links import save_cached,link_person


class ExcelLinksTests(unittest.TestCase):
    def test_link_and_cached_value(self):
        root=Path(__file__).parent/'_runs'/('links_'+uuid.uuid4().hex);root.mkdir(parents=True)
        person=openpyxl.Workbook();person.active.title='个人赋分表'
        person.active['E3']=10;person.active['E4']=5;person.active['E5']='=SUM(E3:E4)'
        save_cached(person,root/'个人.xlsx',{'E5':15})
        book=openpyxl.Workbook();book.active.title='岗位汇总'
        link_person(book,book.active,2,Path('个人.xlsx'),5,15)
        save_cached(book,root/'汇总.xlsx',{'E2':15})
        self.assertEqual(openpyxl.load_workbook(root/'个人.xlsx',data_only=True).active['E5'].value,15)
        self.assertEqual(openpyxl.load_workbook(root/'汇总.xlsx',data_only=True).active['E2'].value,15)
        loaded=openpyxl.load_workbook(root/'汇总.xlsx')
        self.assertEqual(loaded.active['E2'].value,"='[1]个人赋分表'!$E$5")
        self.assertEqual(len(loaded._external_links),1)
        with ZipFile(root/'汇总.xlsx') as z:
            self.assertIn('externalLinkPath',z.read('xl/externalLinks/_rels/externalLink1.xml.rels').decode())


if __name__=='__main__':unittest.main()
