"""验证汇总只同步数值列、保留个人修改、拒绝不明确分数。"""
import sys
import unittest
import uuid
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import openpyxl
from refresh_summary import refresh_summary


class RefreshTests(unittest.TestCase):
    def setUp(self):
        self.root=Path(__file__).parent/'_runs'/('refresh_'+uuid.uuid4().hex)
        folder=self.root/'个人赋分表'/'测试岗位';folder.mkdir(parents=True)
        self.personal=folder/'123_测试.xlsx';self.summary=self.root/'测试岗位.xlsx'
        b=openpyxl.Workbook();s=b.active;s.title='个人赋分表'
        s['E3']=10;s['E4']=0;s['A6']='建议分小计';s['E6']='=SUM(E3:E5)';b.save(self.personal)
        b=openpyxl.Workbook();s=b.active;s.title='岗位汇总'
        s.append(['身份证号','姓名','考生类别','资格建议结论','建议分小计','核验状态','说明'])
        s.append(['123','测试','国内','待复核',88,'人工修改状态','保留说明']);b.save(self.summary)

    def test_only_score_changes(self):
        original=self.personal.read_bytes()
        refresh_summary(self.root)
        b=openpyxl.load_workbook(self.summary)
        self.assertEqual(b.active['E2'].value,10)
        self.assertEqual(b.active['F2'].value,'人工修改状态')
        self.assertEqual(b.active['G2'].value,'保留说明')
        self.assertEqual(original,self.personal.read_bytes())
        self.assertEqual(len(list((self.root/'内部记录').rglob('测试岗位.xlsx'))),1)

    def test_invalid_score_preserves_summary(self):
        b=openpyxl.load_workbook(self.personal);b.active['E3']='=5+5';b.save(self.personal)
        original=self.summary.read_bytes()
        with self.assertRaises(ValueError):refresh_summary(self.root)
        self.assertEqual(original,self.summary.read_bytes())


if __name__=='__main__':unittest.main()
