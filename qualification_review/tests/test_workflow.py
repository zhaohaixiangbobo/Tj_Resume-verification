"""单人合成数据验证导出、人工复核回填和输入变化后的复核失效。"""
import contextlib
import io
import json
from pathlib import Path
import shutil
import sys
import uuid
import unittest
import openpyxl
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from review import run,ROOT

class WorkflowTest(unittest.TestCase):
    def test_review_roundtrip(self):
        # 普通目录继承工作区ACL，兼容Windows受限沙箱。合成文件集中留在tests/_runs。
        with contextlib.nullcontext(Path(__file__).parent/'_runs'/uuid.uuid4().hex) as folder:
            base=Path(folder); app=base/'app';app.mkdir(parents=True)
            cfg=json.loads((ROOT/'config.json').read_text(encoding='utf-8'))
            for name in cfg['standard_sha256']:
                shutil.copyfile(ROOT.parent/name,base/name)
            cfg['personnel_file']='person.xlsx';cfg['jobs_file']='jobs.xlsx';cfg['personnel_sheet']='人员'
            (app/'config.json').write_text(json.dumps(cfg,ensure_ascii=False),encoding='utf-8')
            fields={'姓名':'合成测试人员','身份证号码':'110101200401010010','岗位编码':'001','出生日期':'2004-01-01',
                    '考生类别':'国内暂未毕业','政治面貌':'共青团员','最高学历':'大学本科','最高学位':'学士',
                    '最高学历毕业院校':'燕山大学','最高学历毕业时间':'2027-06-30','最高学历所学专业':'软件工程',
                    '最高学历学习形式':'全日制','学习经历':'燕山大学 2023-09 2027-06 本科 软件工程 全日制 无',
                    '奖学金':'无','表彰奖励':'无','英语等级':'六级(CET-6)425分及以上','计算机等级':'计算机二级',
                    '在读期间有无补考记录':'否','高等教育期间是否取得奖学金':'否','高等教育期间是否获得表彰奖励':'否',
                    '法律职业资格证书，其他':'无'}
            b=openpyxl.Workbook();s=b.active;s.title='人员';s.append(list(fields));s.append(list(fields.values()));b.save(base/'person.xlsx')
            b=openpyxl.Workbook();s=b.active;s.append(['岗位编码','学历要求','本科专业要求','研究生专业要求','其他资格条件要求']);s.append(['001','大学本科','软件工程','','']);b.save(base/'jobs.xlsx')
            with contextlib.redirect_stdout(io.StringIO()):
                first=run(app/'config.json',out_arg=base/'first')
            file=first/'资格审查与赋分试算.xlsx';b=openpyxl.load_workbook(file);s=b['复核录入']
            for row in s.iter_rows(min_row=2):
                if row[3].value=='S13':
                    row[9].value=5;row[10].value='合成测试：按业务标准综合评定适配较高';row[11].value='测试复核人'
            b.save(base/'review.xlsx')
            with contextlib.redirect_stdout(io.StringIO()):
                second=run(app/'config.json',review_path=base/'review.xlsx',out_arg=base/'second')
            b=openpyxl.load_workbook(second/'资格审查与赋分试算.xlsx',data_only=True)
            row=next(r for r in b['逐项依据'].iter_rows(min_row=2,values_only=True) if r[3]=='S13')
            self.assertEqual(row[6],5);self.assertEqual(row[5],'已复核')
            self.assertEqual(b['人员汇总']['E2'].value,'待复核')
            self.assertIsNone(b['人员汇总']['I2'].value)
            b=openpyxl.load_workbook(base/'person.xlsx');b['人员']['D2']='2004-01-02';b.save(base/'person.xlsx')
            with contextlib.redirect_stdout(io.StringIO()),self.assertRaisesRegex(ValueError,'输入已变化'):
                run(app/'config.json',review_path=base/'review.xlsx',out_arg=base/'third')

if __name__=='__main__':unittest.main()
